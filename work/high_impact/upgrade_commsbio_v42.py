#!/usr/bin/env python3
from pathlib import Path
import re,json,pandas as pd
from docx import Document
from docx.shared import Inches,Pt
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'outputs/研究升级/Communications_Biology_v4';V=ROOT/'outputs/研究升级/high_impact_v4/osteoporosis_osteoblast_validation'

def replace_start(d,start,text):
 h=[p for p in d.paragraphs if p.text.startswith(start)]
 if len(h)!=1:raise RuntimeError((start,len(h)))
 h[0].text=text

def insert_before(d,start,blocks):
 h=[p for p in d.paragraphs if p.text==start]
 if not h:h=[p for p in d.paragraphs if p.text.startswith(start)]
 if len(h)!=1:raise RuntimeError((start,len(h)))
 for style,text in blocks:
  p=d.add_paragraph(text,style=style);h[0]._p.addprevious(p._p)
 return h[0]

def add_table(d,headers,rows,font=7):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for j,h in enumerate(headers):t.rows[0].cells[j].text=str(h)
 for row in rows:
  c=t.add_row().cells
  for j,v in enumerate(row):c[j].text=str(v)
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    for r in p.runs:r.font.size=Pt(font)
 return t

def main_doc():
 p=OUT/'Manuscript_Communications_Biology.docx';d=Document(p)
 abstract=('Matrix stiffness regulates mesenchymal stromal cells, but the generalizability and osteoporosis relevance of its transcriptomic effects are unclear. We systematically identified seven public stiffness studies comprising 13 contrasts and 12,241 genes, preserving each GEO series as the inference unit. No gene passed false-discovery-rate correction; study-level correlations ranged from -0.253 to 0.136; and 109 pooled-gene pathway signals fell to zero under study-unit inference. Response magnitude nevertheless localized to human osteoblast-lineage programs. The spatial program remained supported after matching for coverage, uncertainty and gene length (6/7 studies; exact sign-flip P=0.0156). Independent GSE147390 osteoblast markers gave partial support (sign-flip P=0.0391; sign-count P=0.0625; one donor). In 12 independent primary-osteoblast donors, stiffness-response magnitude correlated weakly with osteoporotic-fracture versus osteoarthritis separation (rho=0.0254; exact label-permutation P=0.0303), whereas signed effects and the fixed signature did not transfer. The signature also failed in ITGA11 osteogenic states, mechanical-intervention ATAC-seq and bone-mineral-density genetics. Public data support distributed, lineage-localized mechanosensitivity, while arguing against a universal stiffness gene signature or osteoporosis biomarker.')
 replace_start(d,'Matrix stiffness regulates mesenchymal stromal cell behavior, but',abstract)
 target=insert_before(d,'Discussion',[
  ('Heading 2','Stiffness-response magnitude weakly tracks disease separation in multi-donor osteoblasts'),
  ('Normal','We next analyzed GSE156508, which profiled primary femoral-head osteoblasts from six women with osteoporotic hip fracture and six women undergoing hip replacement for severe osteoarthritis [20]. Among 11,580 genes shared with the stiffness meta-analysis, absolute stiffness meta z scores correlated weakly with absolute fracture-versus-osteoarthritis Welch t statistics (Spearman rho=0.0254). Treating donors as the inference unit, the observed correlation exceeded 96.97% of all 924 balanced label allocations (exact one-sided P=0.0303; Fig. 4a,b). All 12 leave-one-donor-out correlations were positive (range 0.0154-0.0315), although only five deletion analyses had P<0.05.'),
  ('Normal','The association concerned response magnitude rather than a transferable direction. Signed stiffness effects did not correlate with disease effects (rho=0.0518; exact two-sided P=0.361). The fixed 200-gene signature also failed to distinguish fracture from osteoarthritis at the sample level (171 mapped genes; score difference 0.00653; exact two-sided P=0.247; Fig. 4c). Neither the GSE147390 program (Holm P=0.405) nor the spatial osteoblast union (unadjusted P=0.0390; Holm P=0.0779) showed multiplicity-corrected enrichment for disease-effect magnitude. Thus, the clinical cohort supports a small distributed association, not a diagnostic panel.'),
 ])
 pic=d.add_paragraph();pic.alignment=1;pic.add_run().add_picture(str(OUT/'Figure_4.png'),width=Inches(6.65));target._p.addprevious(pic._p)
 cap=d.add_paragraph('Figure 4  Multi-donor primary osteoblast validation. a, Exact null distribution across all 924 balanced assignments of six fracture and six osteoarthritis donors; orange marks the observed correlation. b, Gene-level association between absolute stiffness meta z score and absolute disease Welch t statistic. c, Donor-level fixed signature scores; horizontal bars show group means. OA, osteoarthritis.');cap.runs[0].bold=True;target._p.addprevious(cap._p)
 insert_before(d,'This study is limited',[
  ('Normal','The multi-donor primary osteoblast analysis adds direct skeletal-disease relevance while reinforcing the distributed model. Its statistically unusual combination—a small magnitude correlation with no signed or fixed-signature transfer—is consistent with broad mechanosensitivity across many genes rather than a compact disease classifier. The exact label test avoids treating 11,580 correlated genes as independent replicates, and positive leave-one-donor-out estimates reduce concern about a single influential patient.')])
 replace_start(d,'This study is limited','This study is limited by small source studies, heterogeneous processed data, incomplete reporting of matrix chemistry and biological replication, and only seven independent stiffness series. Moderator analysis was underpowered. Multi-contrast series required aggregation, although a one-contrast-per-series sensitivity agreed. GSE147390 contained one donor. GSE156508 contained 12 donors but compared osteoporotic fracture with severe osteoarthritis rather than healthy bone; fracture consequences, osteoarthritis biology, age or treatment could contribute to the difference. The primary clinical rho was small and only 5/12 leave-one-donor-out tests remained individually significant. The ATAC perturbation contained two samples per group. Single-cell and spatial programs from Chai et al. are distinct datasets within one publication. None of the analyses establish causality or clinical prediction.')
 insert_before(d,'Statistics and reproducibility',[
  ('Heading 2','Multi-donor primary osteoblast disease analysis'),
  ('Normal','GSE156508 normalized GPL16686 expression profiles comprised six osteoporotic-fracture and six severe-osteoarthritis donors [20]. Transcript clusters were mapped with Bioconductor hugene20sttranscriptcluster.db 8.8.0; one-to-many probe mappings were excluded and symbols were collapsed by the median. The primary statistic was Spearman correlation between absolute stiffness meta z and absolute Welch fracture-versus-osteoarthritis t. Its one-sided P value enumerated all 12 choose 6 = 924 balanced donor-label allocations. Secondary signed correlation and sample-level fixed-signature score used exact two-sided label tests. Two prespecified program-magnitude tests used exact one-sided label permutations and Holm correction. Twelve leave-one-donor-out analyses repeated the primary test over each complete allocation space. The comparison was interpreted as fracture-associated osteoporosis versus an osteoarthritis surgical control, not healthy bone.')])
 replace_start(d,'Primary transcriptomic data are available','Primary transcriptomic data are available from NCBI GEO under GSE193021, GSE181512, GSE226411, GSE288678, GSE55867, GSE255574 and GSE310514. Independent osteoblast validations used GSE147390 and the 12-donor primary osteoblast disease cohort GSE156508. External datasets are GSE166824, GSE22011, GSE152708, GSE35958, GSE230665, GSE276529, GSE317531 and GSE287556. Human femoral single-cell data are GSE317069, spatial bone data GSE299207 and the eBMD GWAS GWAS Catalog GCST90726625 [9]. Search exports, screening log, derived data and figure source tables accompany this submission.')
 insert_before(d,'Acknowledgements',[
  ('Normal','20. Panach, L. Expression data in primary osteoblasts obtained from women with osteoporotic fracture or severe osteoarthritis. Gene Expression Omnibus GSE156508 (2020).')])
 d.save(p);print('abstract_words',len(re.findall(r"\b[\w-]+\b",abstract)))

def supplement():
 p=OUT/'Supplementary_Information_Communications_Biology.docx';d=Document(p)
 d.add_heading('Multi-donor primary osteoblast validation',1)
 s=json.loads((V/'osteoporosis_osteoblast_validation_summary.json').read_text())
 rows=[['Primary absolute correlation',s['shared_genes'],f"rho={s['primary']['rho']:.5f}",f"{s['primary']['one_sided_exact_label_permutation_p']:.5f}",'one-sided exact'],['Signed correlation',s['shared_genes'],f"rho={s['secondary_signed']['rho']:.5f}",f"{s['secondary_signed']['two_sided_exact_label_permutation_p']:.5f}",'two-sided exact'],['Fixed 200-gene score',171,f"difference={s['fixed_signature']['fracture_minus_OA_score']:.5f}",f"{s['fixed_signature']['two_sided_exact_label_permutation_p']:.5f}",'two-sided exact']]
 add_table(d,['analysis','n genes','estimate','P','test'],rows,7)
 d.add_paragraph('Supplementary Table 15  GSE156508 primary and fixed-signature analyses. All P values enumerate the complete balanced donor-label allocation space.')
 l=pd.read_csv(V/'leave_one_donor_out_primary.csv')
 add_table(d,['omitted sample','group','rho','exact P'],[[r.omitted_sample,r.omitted_group,f'{r.rho:.5f}',f'{r.exact_p:.5f}'] for r in l.itertuples()],5.6)
 d.add_paragraph('Supplementary Table 16  Leave-one-donor-out audit. All 12 correlations were positive; five exact P values were below 0.05.')
 d.add_paragraph('GSE156508 uses severe osteoarthritis as the comparator. It is therefore a disease-contrast validation, not an osteoporosis-versus-healthy diagnostic study. The small primary effect and null signed/signature results preclude biomarker claims.')
 d.add_heading('Clinical validation quality controls',2)
 add_table(d,['check','result'],[['Unique GPL16686 probes mapped','28,690'],['Collapsed official gene symbols','26,015'],['Genes shared with stiffness meta-analysis','11,580'],['Complete balanced allocations','924'],['Leave-one-donor-out rho positive','12/12'],['Leave-one-donor-out P<0.05','5/12']],6.4)
 d.add_paragraph('Supplementary Table 17  Processing and robustness checks for GSE156508. Probe identifiers with more than one official symbol were excluded before outcome analysis.')
 add_table(d,['program','mapped genes','effect vs complement','exact P','Holm P'],[[z['program'],z['mapped_genes'],f"{z['effect_vs_complement']:.5f}",f"{z['one_sided_exact_p']:.5f}",f"{z['Holm_p']:.5f}"] for z in s['program_disease_magnitude']],6.2)
 d.add_paragraph('Supplementary Table 18  Prespecified osteoblast-program disease-magnitude tests. Neither program passed Holm correction.')
 d.save(p)

def cover():
 p=OUT/'Cover_Letter_Communications_Biology.docx';d=Document(p)
 replace_start(d,'We challenged that interpretation','We challenged that interpretation using stricter gene-length matching, one-representative-contrast analysis, leave-one-study-out checks and a new 12-donor primary human osteoblast cohort. Stiffness-response magnitude showed a small exact-permutation association with osteoporotic-fracture versus osteoarthritis separation, but signed effects and the fixed 200-gene signature did not transfer. Donor-resolved ITGA11 states, bidirectional mechanical-intervention ATAC-seq and eBMD statistics from 448,010 participants further bounded the claim.')
 d.save(p)
if __name__=='__main__':main_doc();supplement();cover()
